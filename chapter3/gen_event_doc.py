import openpyxl as excel, docx, os

# 設定
in_file = "meibo.xlsx"
template_file = "template-event.docx"
save_dir = os.path.join(os.path.dirname(__file__), "events")

# 作成するWordファイルを保存するフォルダを作る
if not os.path.exists(save_dir):
    os.mkdir(save_dir)


# Excelファイルを読んでリストで得る
def read_book():
    result = []
    sheet = excel.load_workbook(in_file).active
    for row in sheet.iter_rows(min_row=2):
        v = [c.value for c in row]
        if v[0] is None:
            break
        result.append(v)
    return result


# 顧客の数だけ案内状を生成
for person in read_book():
    name, zipno, addr = person
    card = {"住所：▲▲▲": addr, "●●様へ": name + "様へ"}
    # Wordテンプレートを読む
    doc = docx.Document(template_file)
    # 内容を書き換える
    for p in doc.paragraphs:
        # テキストを書き換え
        for k, v in card.items():
            if p.text.find(k) >= 0:
                p.text = p.text.replace(k, v)
                p.runs[0].font.bold = True
    # Wordファイルへ保存
    save_file = os.path.join(save_dir, name + "様.docx")
    print("save:", save_file)
    doc.save(save_file)
